"""Read-only extraction of official workbook, preserving cells and scope."""
import json,csv,re,unicodedata,hashlib
from pathlib import Path
import openpyxl
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'data/processed';OUT.mkdir(exist_ok=True)
source=ROOT/'data/collected/references/80_phu-luc-1.xlsx'
book=openpyxl.load_workbook(source,data_only=True,read_only=True)
def norm(s):
 s=unicodedata.normalize('NFD',s.lower().replace('đ','d'));s=''.join(c for c in s if unicodedata.category(c)!='Mn');return re.sub(r'\s+',' ',re.sub(r'^(xa|phuong|dac khu)\s+','',s)).strip()
communes=json.loads((ROOT/'data/boundaries/kon-tum-new-communes-full.geojson').read_text(encoding='utf-8'))['features']
lookup={norm(f['properties']['fullName']):f['properties'] for f in communes}
cols=['sequence','source_commune_name','natural_area_ha','forest_and_immature_plantation_ha','natural_forest_ha','established_plantation_ha','immature_plantation_ha','forest_planning_area_ha','special_use_planning_ha','protection_planning_ha','production_planning_ha','official_forest_cover_percent']
rows=[];allrows=[]
for i,values in enumerate(book['bieu 4'].iter_rows(min_row=11,max_row=106,max_col=12,values_only=True),11):
 if not isinstance(values[0],int):continue
 item=dict(zip(cols,values));item.update(source_sheet='bieu 4',source_row=i,as_of='2025-12-31',decision='1047/QĐ-UBND 31/03/2026',unit='ha except percent',source_url='https://kiemlam.quangngai.gov.vn/items/phuluc/80_phu-luc-1.xlsx')
 item['established_forest_ha']=round((values[4] or 0)+(values[5] or 0),2)
 item['source_missing_numeric_cells']=[cols[j] for j in range(2,12) if values[j] is None]
 item['area_reconciliation_difference_ha']=round(values[3]-sum(values[j] or 0 for j in [4,5,6]),6)
 allrows.append(item)
 if norm(values[1]) in lookup:
  prop=lookup[norm(values[1])];item={**item,'commune_code':prop['code'],'commune_name':prop['fullName'],'match_method':'normalized exact name; official source has no code','scope':'full new commune; not clipped to historical province'};rows.append(item)
assert len(rows)==40,[(k,v['code']) for k,v in lookup.items() if v['code'] not in {r['commune_code'] for r in rows}]
assert len({r['commune_code'] for r in rows})==40
assert all(abs(r['area_reconciliation_difference_ha'])<.01 for r in allrows)
(OUT/'forest-statistics-2025-all-communes.json').write_text(json.dumps(allrows,ensure_ascii=False,indent=2),encoding='utf-8')
(OUT/'forest-statistics-2025-kon-tum-communes.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2),encoding='utf-8')
with (OUT/'forest-statistics-2025-kon-tum-communes.csv').open('w',encoding='utf-8-sig',newline='') as f:
 writer=csv.DictWriter(f,fieldnames=list(rows[0]));writer.writeheader();writer.writerows(rows)
raw={s.title:[list(r) for r in s.iter_rows(max_col=min(s.max_column,20),values_only=True)] for s in book}
(OUT/'official-workbook-extracted.json').write_text(json.dumps(raw,ensure_ascii=False,indent=2),encoding='utf-8')
total=sum(r['natural_area_ha'] for r in rows);forest=sum(r['established_forest_ha'] for r in rows)
report=dict(matched_communes=len(rows),source_total_communes=len(allrows),source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),all_rows_area_reconciled=True,aggregate=dict(natural_area_ha=total,established_forest_ha=forest,derived_weighted_forest_cover_percent=forest/total*100,scope='sum of 40 whole current communes; not official historical-AOI statistic'),notes=['Source column D includes immature plantation; it is NOT established forest area.','Blank source cells preserved as null. Additive reconciliation interprets blanks as zero only for derived checks.','Use published percent as reported, do not silently replace with recalculated percent.'])
(OUT/'official-statistics-validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=True))
